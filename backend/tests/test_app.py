import io
import os
import sys
import time
import tempfile

import pytest

os.environ["VOLT_DB"] = os.path.join(tempfile.mkdtemp(), "test.db")
os.environ["VOLT_SECRET"] = "test-secret-key-test-secret-key-12345"
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app.main import app  # noqa: E402

SAMPLE = os.path.join(os.path.dirname(__file__), "..", "..", "samples", "dpdc_sample.pdf")


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield c


@pytest.fixture(scope="module")
def analysis(client):
    with open(SAMPLE, "rb") as f:
        r = client.post("/api/analyze", files={"file": ("dpdc.pdf", f.read(), "application/pdf")})
    assert r.status_code == 200, r.text
    return r.json()


def test_dpdc_parse_exact(analysis):
    k = analysis["kpis"]
    assert k["n_recharges"] == 19
    assert k["total"] == 38850
    assert round(k["energy"], 2) == 35102.07
    assert round(k["vat"], 2) == 1850.04
    assert k["fixed"] == 2080
    assert analysis["meta"]["utility"] == "DPDC"
    assert analysis["meta"]["customer_no"] == "37001602"
    assert analysis["meta"]["meter_no"] == "020120093549"
    assert analysis["meta"]["confidence"] == "high"
    assert analysis["meta"]["warnings"] == []


def test_every_row_reconciles(analysis):
    for r in analysis["records"]:
        s = r["service"] + r["demand"] + r["rent"] + r["arrear"] + r["rebate"] + r["vat"] + r["energy"]
        assert abs(r["amount"] - s) < 0.02, r


def test_seasons(analysis):
    s = {x["season"]: x for x in analysis["seasons"]}
    assert s["Winter"]["per_30"] == 2273 and s["Winter"]["days"] == 59
    assert s["Summer"]["per_30"] == 3571 and s["Summer"]["days"] == 92
    assert s["Rainy"]["per_30"] == 5270 and s["Rainy"]["days"] == 122
    assert s["Winter"]["index"] == 1.0 and s["Rainy"]["index"] == 2.32
    assert sum(x["total"] for x in analysis["seasons"]) == 38850 - 2000  # partial October excluded


def test_cycles_merge_same_day(analysis):
    dates = [c["date"] for c in analysis["cycles"]]
    assert len(dates) == len(set(dates)) == 18
    aug16 = next(c for c in analysis["cycles"] if c["date"] == "2026-08-16")
    assert aug16["amount"] == 2000 and aug16["days"] == 10
    assert analysis["cycles"][-1]["closed"] is False


def test_speed(client):
    with open(SAMPLE, "rb") as f:
        data = f.read()
    t = time.time()
    r = client.post("/api/analyze", files={"file": ("dpdc.pdf", data, "application/pdf")})
    assert r.status_code == 200
    assert time.time() - t < 5


def test_export_xlsx(client, analysis):
    body = {"meta": analysis["meta"], "records": analysis["records"], "season_map": analysis["season_map"]}
    r = client.post("/api/export", json=body)
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    assert "attachment" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "Seasonal", "Monthly", "Cycles", "Ledger", "Assumptions", "Charts"]
    ws = wb["Ledger"]
    assert ws["E23"].value == "=SUM(E4:E22)"  # formulas, not hardcoded results
    out = os.environ.get("VOLT_TEST_XLSX_OUT")
    if out:
        open(out, "wb").write(r.content)


def test_reject_bad_inputs(client):
    r = client.post("/api/analyze", files={"file": ("x.pdf", b"not a pdf", "application/pdf")})
    assert r.status_code == 422 and "PDF" in r.json()["detail"]
    r = client.post("/api/analyze", files={"file": ("x.exe", b"MZ....", "application/octet-stream")})
    assert r.status_code == 422
    r = client.post("/api/analyze", files={"file": ("e.csv", b"", "text/csv")})
    assert r.status_code == 422
    r = client.post("/api/export", json={"meta": {}, "records": []})
    assert r.status_code == 422


def test_generic_csv_desco_style(client):
    csv = (
        "Recharge Date,Recharge Amount,Energy Cost,VAT,Demand Charge,Meter Rent\n"
        "15/01/2026,1000,880,43,60,17\n"
        "10/02/2026,1000,880,43,60,17\n"
        "05/05/2026,2000,1800,90,60,50\n"
    )
    r = client.post("/api/analyze", files={"file": ("desco.csv", csv.encode(), "text/csv")})
    assert r.status_code == 200, r.text
    j = r.json()
    assert j["kpis"]["n_recharges"] == 3 and j["kpis"]["total"] == 4000
    assert j["meta"]["confidence"] in ("medium", "low")


def test_auth_and_saved_reports(client, analysis):
    body = {"meta": analysis["meta"], "records": analysis["records"], "season_map": analysis["season_map"], "filename": "dpdc.pdf"}
    assert client.get("/api/reports").status_code == 401
    assert client.post("/api/reports", json=body).status_code == 401

    r = client.post("/api/auth/register", json={"email": "a@example.com", "password": "password123"})
    assert r.status_code == 200
    h = {"Authorization": "Bearer " + r.json()["token"]}
    assert client.post("/api/auth/register", json={"email": "A@example.com", "password": "password123"}).status_code == 409
    assert client.post("/api/auth/login", json={"email": "a@example.com", "password": "wrongpass1"}).status_code == 401
    assert client.post("/api/auth/register", json={"email": "bad", "password": "password123"}).status_code == 422
    assert client.post("/api/auth/register", json={"email": "c@example.com", "password": "short"}).status_code == 422

    rid = client.post("/api/reports", json=body, headers=h).json()["id"]
    lst = client.get("/api/reports", headers=h).json()
    assert len(lst) == 1 and lst[0]["id"] == rid
    full = client.get(f"/api/reports/{rid}", headers=h).json()
    assert full["kpis"]["total"] == 38850
    assert client.get(f"/api/reports/{rid}/xlsx", headers=h).status_code == 200

    # another user cannot see it
    r2 = client.post("/api/auth/register", json={"email": "b@example.com", "password": "password123"})
    h2 = {"Authorization": "Bearer " + r2.json()["token"]}
    assert client.get(f"/api/reports/{rid}", headers=h2).status_code == 404
    assert client.delete(f"/api/reports/{rid}", headers=h2).status_code == 404
    assert client.get("/api/reports", headers=h2).json() == []

    assert client.delete(f"/api/reports/{rid}", headers=h).status_code == 200
    assert client.get(f"/api/reports/{rid}", headers=h).status_code == 404


def test_frontend_served(client):
    assert "VoltLedger" in client.get("/").text
    assert client.get("/static/app.js").status_code == 200
    assert "default-src 'self'" in client.get("/").headers["content-security-policy"]
